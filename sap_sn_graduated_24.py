import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

def hdr(ws, row, col, value, bg="1F4E79", fg="FFFFFF", bold=True, size=11, wrap=True):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = PatternFill("solid", fgColor=bg)
    c.font = Font(color=fg, bold=bold, size=size, name="Calibri")
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    thin = Side(style="thin", color="AAAAAA")
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    return c

def cell(ws, row, col, value, bg=None, bold=False, wrap=True, align="left", size=10):
    c = ws.cell(row=row, column=col, value=value)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    c.font = Font(bold=bold, size=size, name="Calibri")
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    thin = Side(style="thin", color="DDDDDD")
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    return c

def sec(ws, row, col, span, title, bg="2E75B6"):
    c = ws.cell(row=row, column=col, value=title)
    c.fill = PatternFill("solid", fgColor=bg)
    c.font = Font(color="FFFFFF", bold=True, size=12, name="Calibri")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+span-1)

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 1 — สรุปตัวเลือกที่ใช้ได้จริง (อายุ 24 จบแล้ว)
# ══════════════════════════════════════════════════════════════════════════════
ws1 = wb.active
ws1.title = "1_ตัวเลือกที่ใช้ได้ (24 จบแล้ว)"
ws1.sheet_view.showGridLines = False

sec(ws1, 1, 1, 7, "  ✅  ตัวเลือกที่ใช้ได้จริง — อายุ 24 จบการศึกษาแล้ว (ไม่ใช่นักศึกษา)")

headers = ["เส้นทาง", "ประเทศ", "วีซ่า", "ตำแหน่งที่หาได้", "เงินเดือน", "ใช้ได้?", "หมายเหตุ"]
for i, h in enumerate(headers, 1):
    hdr(ws1, 2, i, h)

data = [
    ["🅐 Work & Holiday\n→ หางาน Junior\nSAP ABAP / ServiceNow\n[แนะนำสุด]",
     "🇦🇺 ออสเตรเลีย",
     "WHV Subclass 462\n(อายุ 18–30 ✅\nจบ ป.ตรี ✅)",
     "• Junior SAP ABAP Developer\n• SAP Consultant (entry)\n• ServiceNow Junior Dev\n• IT Support (SAP)\nที่: Deloitte, KPMG, Accenture,\nInfosys, Wipro, TCS, SAP AU",
     "AUD 65,000–100,000/ปี\n(SAP ABAP)\nAUD 70,000–114,000/ปี\n(ServiceNow)",
     "✅ ใช้ได้\n100%",
     "ทำงานกับนายจ้างเดียว\nได้ไม่เกิน 6 เดือน\n→ ขอให้ sponsor\nSubclass 482 ต่อ\n\n⚠️ เปิดรับปี 2026\n(ปี 2025 ปิดแล้ว)"],
    ["🅑 Work & Holiday\n→ นิวซีแลนด์",
     "🇳🇿 นิวซีแลนด์",
     "WHV\n(อายุ 18–30 ✅\nQuota 100/ปี ⚠️)",
     "• SAP Junior Developer\n• ServiceNow Admin\n• IT Consultant\n(ตลาดเล็กกว่า AU)",
     "NZD 65,000–110,000/ปี",
     "✅ ใช้ได้\nแต่ quota น้อย",
     "Quota 100 คน/ปีเท่านั้น\nต้องสมัครเร็วมาก\nแนะนำ: ลอง AU ก่อน"],
    ["🅒 Graduate Program\nที่ Consulting Firms\n(ออสเตรเลีย)",
     "🇦🇺 ออสเตรเลีย",
     "WHV ก่อน\n→ Employer Sponsor\n→ Subclass 482",
     "• EY Graduate – ServiceNow\n• Accenture Graduate\n  SAP Practice\n• Deloitte Graduate\n  Technology\n• KPMG Ignition",
     "AUD 65,000–80,000/ปี\n(Grad year 1–2)",
     "✅ ใช้ได้\n(recent grad ✅)",
     "หลายบริษัทรับ\n'recent graduate'\nไม่ต้องเป็น student\nสมัครผ่าน WHV แล้ว\nขอ sponsor ต่อ"],
    ["🅓 Skilled Worker Visa\n→ เยอรมนี / ยุโรป\n[ถ้ามีประสบการณ์ 1 ปีขึ้น]",
     "🇩🇪 เยอรมนี\n🇳🇱 เนเธอร์แลนด์\n🇬🇧 UK",
     "Skilled Worker Visa\n(Germany)\nGlobal Talent Visa (UK)",
     "• SAP ABAP Developer\n• ServiceNow Developer\n• SAP Consultant\n• ITSM Specialist\nที่: SAP SE, Accenture DE,\nCapgemini, T-Systems",
     "SAP ABAP:\n€42,000–65,000/ปี\nServiceNow:\n€46,000–63,000/ปี",
     "✅ ใช้ได้\n(มีประสบการณ์\n1 ปีขึ้น)",
     "Germany Immigration Act 2025\nง่ายขึ้นมากสำหรับ non-EU IT\nขาดคน 109,000 ตำแหน่ง\n\n💡 เก็บ exp ในไทย 1 ปีก่อน\nแล้ว apply Germany"],
    ["🅔 Employment Pass\n→ สิงคโปร์",
     "🇸🇬 สิงคโปร์",
     "Employment Pass\n(ต้องมี job offer\n+ เงินเดือน\nSGD 5,000+/เดือน)",
     "• ServiceNow Developer\n• SAP Consultant\n• ITSM Consultant\nที่: ServiceNow HQ,\nAccenture SG, Deloitte SG",
     "SGD 60,000–100,000/ปี\n(ServiceNow)\nSGD 60,000–90,000/ปี\n(SAP)",
     "✅ ใช้ได้\n(แต่ต้องมี\njob offer ก่อน)",
     "ไม่มี WHV กับไทย\nต้อง apply job\nจากไทยก่อน\nได้ offer → apply EP\n\nServiceNow demand\nสูงมากใน SG ปี 2025"],
    # Crossed out options
    ["❌ SAP iXp Internship\n(Walldorf, Germany)",
     "🇩🇪 เยอรมนี",
     "Student Visa",
     "iXp Intern – ABAP Developer",
     "~€800–1,200/เดือน",
     "❌ ใช้ไม่ได้\n(ต้องเป็น enrolled\nstudent เท่านั้น)",
     "ต้องมีสถานะ\nnกศ. อยู่จึงจะสมัครได้\nจบแล้วสมัครไม่ได้"],
    ["❌ Accenture\nSAP ABAP Internship\n(Poland)",
     "🇵🇱 โปแลนด์",
     "Student Visa",
     "SAP ABAP Intern",
     "~€800–1,150/เดือน",
     "❌ ใช้ไม่ได้\n(ต้องเป็น enrolled\nstudent เท่านั้น)",
     "ระบุชัดว่า\n'enrolled in degree\nprogram'\nจบแล้วสมัครไม่ได้"],
    ["❌ ServiceNow\nOfficial Internship\n(Singapore / USA)",
     "🇸🇬 สิงคโปร์\n🇺🇸 สหรัฐฯ",
     "Student Visa",
     "Software Engineer Intern",
     "$53/hr (US)\nSGD 1,500–3,500/เดือน",
     "❌ ใช้ไม่ได้\n(ต้องเป็น enrolled\nstudent เท่านั้น)",
     "ต้องเป็นนักศึกษา\nที่ยังเรียนอยู่\nเท่านั้น"],
]

row_colors_map = {
    "✅": "E8F9F2",
    "❌": "FFF0F0",
}
for i, row_data in enumerate(data):
    r = i + 3
    status = row_data[5]
    bg = "FFF0F0" if "❌" in status else "E8F9F2"
    if i % 2 == 1 and "❌" not in status:
        bg = "F2FBF6"
    for j, val in enumerate(row_data, 1):
        c = cell(ws1, r, j, val, bg=bg, wrap=True)
        if j == 1:
            c.font = Font(bold=True, size=10, name="Calibri")
            if "❌" in val:
                c.font = Font(bold=True, size=10, name="Calibri", color="999999")
        if j == 6:
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if "❌" in val:
                c.font = Font(bold=True, size=11, color="CC0000", name="Calibri")
            else:
                c.font = Font(bold=True, size=11, color="006600", name="Calibri")
    ws1.row_dimensions[r].height = 90

col_widths = [24, 16, 22, 36, 24, 16, 30]
for i, w in enumerate(col_widths, 1):
    ws1.column_dimensions[get_column_letter(i)].width = w
ws1.row_dimensions[1].height = 30
ws1.row_dimensions[2].height = 35

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 2 — Action Plan ที่เหมาะสมสุด
# ══════════════════════════════════════════════════════════════════════════════
ws2 = wb.create_sheet("2_Action Plan อายุ 24 จบแล้ว")
ws2.sheet_view.showGridLines = False

sec(ws2, 1, 1, 4, "  🗺️  Action Plan — อายุ 24 จบการศึกษาแล้ว (2 เส้นทางหลักที่แนะนำ)")

hdr(ws2, 2, 1, "เส้นทาง")
hdr(ws2, 2, 2, "Timeline (เดือน)")
hdr(ws2, 2, 3, "ประเทศ / วีซ่า")
hdr(ws2, 2, 4, "งบประมาณเตรียมตัว")

plans = [
    ["⭐ เส้นทางที่ 1\nServiceNow\n→ WHV ออสเตรเลีย\n[แนะนำที่สุด]\n(เร็ว + ถูก + หางานง่าย)",
     "เดือน 1–2:\n• เรียน ServiceNow ฟรีที่ developer.servicenow.com\n• สร้าง Personal Dev Instance (ฟรี)\n• ทำ portfolio 2–3 projects\n\nเดือน 3:\n• สอบ ServiceNow CSA (~$250)\n\nเดือน 4–5:\n• สอบ IELTS ให้ได้ 4.5+\n• เตรียมเอกสาร WHV (ปริญญาตรี + bank stmt)\n• สมัคร WHV Subclass 462 ปี 2026\n\nเดือน 6–7:\n• ได้วีซ่า + เดินทางออสเตรเลีย\n\nเดือน 7–12:\n• หางาน Junior ServiceNow Dev\n  ที่ EY AU / Deloitte AU / Accenture AU\n• เงินเดือน AUD 70,000–90,000/ปี\n\nปีที่ 2:\n• ขอ Employer Sponsor → Subclass 482",
     "🇦🇺 ออสเตรเลีย\nWHV Subclass 462\nอายุ 24 ✅ จบ ป.ตรี ✅",
     "ServiceNow training: ฟรี ✅\nCSA Cert: ~8,500 บาท\nIELTS: ~7,000 บาท\nค่าวีซ่า WHV: ~9,500 บาท\nตั๋วเครื่องบิน: ~15,000 บาท\nค่าครองชีพ 1 เดือนแรก: ~50,000 บาท\n─────────────────\nรวมเตรียมตัว: ~90,000 บาท"],
    ["เส้นทางที่ 2\nSAP ABAP\n→ WHV ออสเตรเลีย\n(ใช้เวลาเตรียมมากกว่า\nแต่เงินเดือนสูงกว่า)",
     "เดือน 1–3:\n• เรียน SAP ABAP บน SAP Learning Hub\n  หรือ Udemy / YouTube\n• ใช้ SAP BTP free tier ฝึก\n\nเดือน 4–5:\n• สอบ SAP Certified Associate ABAP\n  (~$550 USD)\n\nเดือน 5–6:\n• สอบ IELTS 4.5+\n• สมัคร WHV 2026\n\nเดือน 7–8:\n• เดินทางออสเตรเลีย\n\nเดือน 8–12:\n• หางาน Junior SAP ABAP Dev\n  ที่ Infosys AU / TCS AU / Wipro AU\n  / SAP Australia / Capgemini AU\n• เงินเดือน AUD 80,000–100,000/ปี\n\nปีที่ 2:\n• ขอ Employer Sponsor → Subclass 482",
     "🇦🇺 ออสเตรเลีย\nWHV Subclass 462\nอายุ 24 ✅ จบ ป.ตรี ✅",
     "SAP ABAP course (Udemy): ~1,500 บาท\nSAP Cert ABAP: ~20,000 บาท\nIELTS: ~7,000 บาท\nค่าวีซ่า WHV: ~9,500 บาท\nตั๋วเครื่องบิน: ~15,000 บาท\nค่าครองชีพ 1 เดือนแรก: ~50,000 บาท\n─────────────────\nรวมเตรียมตัว: ~103,000 บาท"],
    ["เส้นทางที่ 3\nเก็บ Exp ในไทย\n→ Skilled Worker\nเยอรมนี / ยุโรป\n(ถ้าอยากไป Europe)",
     "ปีที่ 1 (ในไทย):\n• ทำงาน SAP ABAP หรือ ServiceNow\n  ในบริษัทไทยก่อน 1 ปี\n• สอบ Certification ให้ได้ CSA/ABAP\n• เรียนภาษาอังกฤษให้ดีขึ้น (B2)\n\nปีที่ 2:\n• สมัครงานบน LinkedIn / germantechjobs.de\n  สำหรับตำแหน่ง Germany/EU\n• ได้ job offer → นายจ้าง sponsor\n  Skilled Worker Visa\n• ยื่นวีซ่าที่สถานทูตเยอรมัน กรุงเทพฯ\n• Processing ~4–12 สัปดาห์\n\nปีที่ 2–3:\n• ทำงานที่เยอรมนี\n  SAP ABAP: €42k–65k/ปี\n  ServiceNow: €46k–63k/ปี",
     "🇩🇪 เยอรมนี\nSkilled Worker Visa\n(ค่าวีซ่า €100 เท่านั้น\nnายจ้าง sponsor)",
     "ค่าวีซ่า: ~3,800 บาท\nตั๋วเครื่องบิน: ~25,000 บาท\nค่าครองชีพเดือนแรก: ~60,000 บาท\nCert (ถ้ายังไม่มี): ~8,500–20,000 บาท\n─────────────────\nรวม: ~90,000–110,000 บาท\n(นายจ้างช่วย relocation บ่อยมาก)"],
]

plan_colors = ["E8F9F2", "E8F4FD", "FFF3E0"]
for i, row_data in enumerate(plans):
    r = i + 3
    ws2.row_dimensions[r].height = 170
    for j, val in enumerate(row_data, 1):
        c = cell(ws2, r, j, val, bg=plan_colors[i], wrap=True)
        if j == 1:
            c.font = Font(bold=True, size=10, name="Calibri")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws2.column_dimensions["A"].width = 22
ws2.column_dimensions["B"].width = 52
ws2.column_dimensions["C"].width = 22
ws2.column_dimensions["D"].width = 34
ws2.row_dimensions[1].height = 30
ws2.row_dimensions[2].height = 35

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 3 — เตรียมตัว ServiceNow ฟรี (แนะนำสุด)
# ══════════════════════════════════════════════════════════════════════════════
ws3 = wb.create_sheet("3_เตรียมตัว ServiceNow (ฟรี)")
ws3.sheet_view.showGridLines = False

sec(ws3, 1, 1, 4, "  📚  เตรียมตัว ServiceNow ฟรีทั้งหมด — ก่อนสมัคร WHV ออสเตรเลีย")

hdr(ws3, 2, 1, "ขั้นตอน")
hdr(ws3, 2, 2, "สิ่งที่ต้องทำ")
hdr(ws3, 2, 3, "แหล่งข้อมูล / ลิงก์")
hdr(ws3, 2, 4, "ค่าใช้จ่าย")

prep_data = [
    ["1️⃣ สร้าง Account\nServiceNow Developer",
     "• ไปที่ developer.servicenow.com\n• สมัครฟรี → ได้ Personal Dev Instance (PDI)\n• PDI คือ ServiceNow จริงๆ ให้ใช้ฟรี\n• Activate instance ใช้ฝึกได้เลย",
     "developer.servicenow.com",
     "ฟรี ✅"],
    ["2️⃣ เรียน Fundamentals\nบน NowLearning",
     "• เรียน 'ServiceNow Fundamentals'\n• เรียน 'Scripting in ServiceNow'\n• เรียน 'Flow Designer'\n• เรียน 'ITSM Implementation'\n→ ครบทุกอย่างที่ต้องรู้สำหรับงาน",
     "nowlearning.servicenow.com\n(ฟรีทั้งหมด!)",
     "ฟรี ✅"],
    ["3️⃣ ทำ Portfolio\nบน PDI",
     "• สร้าง custom application บน PDI\n• Automate workflow ด้วย Flow Designer\n• ทำ integration กับ REST API ภายนอก\n• สร้าง Dashboard + Report\n→ เก็บ screenshots ทำ portfolio",
     "ใช้ PDI ที่ได้จากขั้นตอน 1\ndeveloper.servicenow.com/dev",
     "ฟรี ✅"],
    ["4️⃣ สอบ CSA\n(Certified System\nAdministrator)",
     "• เรียน 'Certified System Administrator'\ncourse บน NowLearning ก่อน\n• ทำ Practice Exam ให้ผ่าน 70%+\n• สอบ online ที่บ้านได้เลย (Proctored)\n• ใช้เวลา 90 นาที 60 ข้อ\n→ ผ่าน 70% = ได้ CSA",
     "nowlearning.servicenow.com\n→ Certifications → CSA\n\nPractice exam ฟรีบน\nnowlearning.servicenow.com",
     "~$250 USD\n(~8,750 บาท)\n\nสอบที่บ้านได้!"],
    ["5️⃣ สอบ CIS-ITSM\n(Certified Implementation\nSpecialist — ITSM)\n[Optional แต่ช่วยมาก]",
     "• ได้ CSA แล้วค่อยสอบ CIS\n• เพิ่มมูลค่าเงินเดือน 20–40%\n• เหมาะกับงาน ITSM Implementation\n  ที่ consulting firms (EY/Deloitte/Accenture)",
     "nowlearning.servicenow.com\n→ Certifications → CIS-ITSM",
     "~$350 USD\n(~12,250 บาท)\n\nสอบหลังมี CSA\nแล้วค่อยทำ"],
    ["6️⃣ หางาน Junior\nServiceNow\nในออสเตรเลีย",
     "• ใช้ WHV Subclass 462\n• สมัครที่:\n  - EY Australia ServiceNow Graduate\n  - Accenture Australia Graduate\n  - Deloitte Technology Graduate\n  - Seek.com.au → ServiceNow\n  - LinkedIn → ServiceNow junior Australia\n• เตรียม CV + Portfolio + CSA cert",
     "ey.com/en_au/careers/\nservicenow-student-opportunities\n\nau.seek.com → ServiceNow\nlinkedin.com → filter Australia",
     "ฟรี\n(ค่าสมัครงาน\nไม่มี)"],
]

prep_colors = ["F2FBF6", "E8F9F2"]
for i, row_data in enumerate(prep_data):
    r = i + 3
    bg = prep_colors[i % 2]
    ws3.row_dimensions[r].height = 75
    for j, val in enumerate(row_data, 1):
        c = cell(ws3, r, j, val, bg=bg, wrap=True)
        if j == 1:
            c.font = Font(bold=True, size=11, name="Calibri")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if j == 4:
            col = "006600" if "ฟรี" in val else "333333"
            bold = "ฟรี" in val
            c.font = Font(bold=bold, size=10, color=col, name="Calibri")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws3.column_dimensions["A"].width = 20
ws3.column_dimensions["B"].width = 50
ws3.column_dimensions["C"].width = 36
ws3.column_dimensions["D"].width = 18
ws3.row_dimensions[1].height = 30
ws3.row_dimensions[2].height = 35

# ══════════════════════════════════════════════════════════════════════════════
# SHEET 4 — บริษัทที่สมัครได้ (ออสเตรเลีย) ผ่าน WHV
# ══════════════════════════════════════════════════════════════════════════════
ws4 = wb.create_sheet("4_บริษัทเป้าหมาย AU WHV")
ws4.sheet_view.showGridLines = False

sec(ws4, 1, 1, 6, "  🏢  บริษัทที่สมัครได้ด้วย WHV ออสเตรเลีย — SAP ABAP & ServiceNow")

hdr(ws4, 2, 1, "บริษัท")
hdr(ws4, 2, 2, "สาย")
hdr(ws4, 2, 3, "ตำแหน่งที่หาได้")
hdr(ws4, 2, 4, "เมือง")
hdr(ws4, 2, 5, "เงินเดือน (ปีแรก)")
hdr(ws4, 2, 6, "หมายเหตุ")

companies = [
    # ServiceNow
    ["EY Australia", "ServiceNow", "Graduate Technology\n(ServiceNow Practice)\n/ Junior Developer", "Sydney\nMelbourne", "AUD 65,000–80,000/ปี",
     "มี ServiceNow student\nprogram โดยตรง\ney.com/en_au/careers"],
    ["Accenture\nAustralia", "ServiceNow\n+ SAP ABAP", "Graduate Program –\nTechnology\n(SAP / ServiceNow)", "Sydney\nMelbourne\nBrisbane", "AUD 68,000–85,000/ปี",
     "รับ recent graduate\nทั้ง SAP และ ServiceNow\naccenture.com/au-en/careers"],
    ["Deloitte\nAustralia", "ServiceNow\n+ SAP ABAP", "Graduate Analyst –\nTechnology Consulting", "Sydney\nMelbourne\nPerth", "AUD 65,000–80,000/ปี",
     "Graduate program\nรับ recent grads\ndeloitte.com/au/careers"],
    ["KPMG\nAustralia", "ServiceNow\n+ SAP", "Graduate –\nTechnology Advisory\n/ Digital Enablement", "Sydney\nMelbourne", "AUD 62,000–75,000/ปี",
     "KPMG Ignition program\nkpmg.com.au/careers"],
    ["PwC Australia", "ServiceNow", "Graduate –\nCloud & Digital\nTransformation", "Sydney\nMelbourne", "AUD 65,000–78,000/ปี",
     "pwc.com.au/careers\nGraduate program"],
    # SAP ABAP
    ["SAP Australia", "SAP ABAP", "Junior ABAP Developer\n/ Associate Consultant", "Sydney\nMelbourne", "AUD 75,000–95,000/ปี",
     "jobs.sap.com\nFilter: Australia\nVisa sponsor พร้อม"],
    ["Infosys\nAustralia", "SAP ABAP", "SAP ABAP Developer\n/ Junior Consultant", "Sydney\nMelbourne\nBrisbane", "AUD 70,000–90,000/ปี",
     "infosys.com/careers\nรับ international\nทำงานใน project ไทย/AU"],
    ["TCS Australia\n(Tata Consultancy\nServices)", "SAP ABAP", "SAP Junior Developer\n/ Associate", "Sydney\nMelbourne", "AUD 68,000–85,000/ปี",
     "tcs.com/careers\nบริษัทอินเดีย\nรับ international มาก"],
    ["Wipro\nAustralia", "SAP ABAP", "SAP ABAP Developer", "Sydney\nMelbourne", "AUD 65,000–85,000/ปี",
     "wipro.com/careers\nรับ WHV workers\nบ่อยมาก"],
    ["Capgemini\nAustralia", "SAP ABAP\n+ ServiceNow", "Junior SAP Consultant\n/ ServiceNow Dev", "Sydney\nMelbourne", "AUD 68,000–88,000/ปี",
     "capgemini.com/au-en/careers\nทั้งสองสาย"],
]

abap_bg = "E8F4FD"
sn_bg = "FFF3E0"
both_bg = "EEF2F7"

for i, row in enumerate(companies):
    r = i + 3
    ws4.row_dimensions[r].height = 60
    sai = row[1]
    if "SAP ABAP" in sai and "ServiceNow" in sai:
        bg = both_bg
    elif "ServiceNow" in sai:
        bg = sn_bg
    else:
        bg = abap_bg
    for j, val in enumerate(row, 1):
        c = cell(ws4, r, j, val, bg=bg, wrap=True)
        if j == 1:
            c.font = Font(bold=True, size=10, name="Calibri")

ws4.column_dimensions["A"].width = 18
ws4.column_dimensions["B"].width = 16
ws4.column_dimensions["C"].width = 26
ws4.column_dimensions["D"].width = 16
ws4.column_dimensions["E"].width = 22
ws4.column_dimensions["F"].width = 30
ws4.row_dimensions[1].height = 30
ws4.row_dimensions[2].height = 35

# Save
out = "/home/user/fefa/SAP_SN_Graduated_24yo.xlsx"
wb.save(out)
print(f"Saved: {out}")
